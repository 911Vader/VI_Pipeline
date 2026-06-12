"""Tests for reporting/spreadsheet.py"""
import pytest

pytest.importorskip("openpyxl")

import openpyxl
from pathlib import Path

from analysis.parser import Variant
from reporting.spreadsheet import write_spreadsheet, MANDATORY_DISCLAIMER, _ACMG_COLOURS, VARIANT_COLUMNS


def _make_variant(
    chrom="chr1", pos=100, ref="A", alt="T",
    gene_symbol="GENE1", acmg_classification="VUS",
    excluded=False, exclusion_reason=None,
    acmg_evidence_codes=None,
    zygosity="heterozygous",
    gnomad_af=0.0001,
):
    v = Variant(
        chrom=chrom, pos=pos, ref=ref, alt=alt,
        gene_symbol=gene_symbol,
        acmg_classification=acmg_classification,
        excluded=excluded,
        exclusion_reason=exclusion_reason,
        acmg_evidence_codes=acmg_evidence_codes or [],
        zygosity=zygosity,
        gnomad_af=gnomad_af,
    )
    return v


# ---------------------------------------------------------------------------
# Basic file creation
# ---------------------------------------------------------------------------

class TestWriteSpreadsheet:
    def test_creates_file(self, tmp_path):
        out = str(tmp_path / "out.xlsx")
        result = write_spreadsheet([_make_variant()], out, {"sample_id": "S1"})
        assert Path(result).exists()

    def test_returns_path(self, tmp_path):
        out = str(tmp_path / "out.xlsx")
        result = write_spreadsheet([_make_variant()], out, {"sample_id": "S1"})
        assert result == out

    def test_creates_parent_dirs(self, tmp_path):
        out = str(tmp_path / "deep" / "nested" / "out.xlsx")
        write_spreadsheet([_make_variant()], out, {"sample_id": "S1"})
        assert Path(out).exists()

    def test_valid_xlsx(self, tmp_path):
        out = str(tmp_path / "out.xlsx")
        write_spreadsheet([_make_variant()], out, {"sample_id": "S1"})
        wb = openpyxl.load_workbook(out)
        assert wb is not None


# ---------------------------------------------------------------------------
# Sheet structure
# ---------------------------------------------------------------------------

class TestSheetStructure:
    EXPECTED_SHEETS = {
        "Info", "Priority Review", "VUS", "In Gene List",
        "High Impact", "Deprioritized", "Excluded",
    }

    def test_correct_sheets_created(self, tmp_path):
        out = str(tmp_path / "out.xlsx")
        write_spreadsheet([_make_variant()], out, {"sample_id": "S1"})
        wb = openpyxl.load_workbook(out)
        assert set(wb.sheetnames) == self.EXPECTED_SHEETS

    def test_info_sheet_first(self, tmp_path):
        out = str(tmp_path / "out.xlsx")
        write_spreadsheet([_make_variant()], out, {"sample_id": "S1"})
        wb = openpyxl.load_workbook(out)
        assert wb.sheetnames[0] == "Info"

    def test_disclaimer_on_info_sheet(self, tmp_path):
        out = str(tmp_path / "out.xlsx")
        write_spreadsheet([_make_variant()], out, {"sample_id": "S1"})
        wb = openpyxl.load_workbook(out)
        ws = wb["Info"]
        all_values = [ws.cell(r, 2).value for r in range(1, ws.max_row + 1)]
        assert MANDATORY_DISCLAIMER in all_values

    def test_vus_sheet_has_header(self, tmp_path):
        out = str(tmp_path / "out.xlsx")
        write_spreadsheet([_make_variant()], out, {"sample_id": "S1"})
        wb = openpyxl.load_workbook(out)
        ws = wb["VUS"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "Gene" in headers
        assert "ACMG Classification" in headers
        assert "Sample ID" in headers

    def test_vus_variant_goes_to_vus_sheet(self, tmp_path):
        """A VUS variant should appear in the VUS sheet."""
        out = str(tmp_path / "out.xlsx")
        write_spreadsheet([_make_variant(acmg_classification="VUS")], out, {"sample_id": "S1"})
        wb = openpyxl.load_workbook(out)
        ws = wb["VUS"]
        # Row 1 = header, row 2 = first variant
        assert ws.max_row == 2

    def test_excluded_variant_goes_to_excluded_sheet(self, tmp_path):
        retained = _make_variant()
        excluded = _make_variant(excluded=True, exclusion_reason="gnomAD AF > 0.01")
        out = str(tmp_path / "out.xlsx")
        write_spreadsheet([retained, excluded], out, {"sample_id": "S1"})
        wb = openpyxl.load_workbook(out)
        ws_excl = wb["Excluded"]
        assert ws_excl.max_row == 2  # header + 1 excluded

    def test_empty_variants_still_creates_sheets(self, tmp_path):
        out = str(tmp_path / "out.xlsx")
        write_spreadsheet([], out, {"sample_id": "S1"})
        wb = openpyxl.load_workbook(out)
        assert "Priority Review" in wb.sheetnames
        assert "Excluded" in wb.sheetnames
        assert "VUS" in wb.sheetnames


# ---------------------------------------------------------------------------
# Data correctness
# ---------------------------------------------------------------------------

class TestDataCorrectness:
    def test_gene_symbol_in_row(self, tmp_path):
        v = _make_variant(gene_symbol="BRCA2", acmg_classification="VUS")
        out = str(tmp_path / "out.xlsx")
        write_spreadsheet([v], out, {"sample_id": "S1"})
        wb = openpyxl.load_workbook(out)
        ws = wb["VUS"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        gene_col = headers.index("Gene") + 1
        assert ws.cell(2, gene_col).value == "BRCA2"

    def test_first_row_has_gene_symbol(self, tmp_path):
        v = _make_variant(gene_symbol="BRCA1", acmg_classification="VUS")
        out = str(tmp_path / "out.xlsx")
        write_spreadsheet([v], out, {"sample_id": "S1"})
        wb = openpyxl.load_workbook(out)
        ws = wb["VUS"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        gene_col = headers.index("Gene") + 1
        assert ws.cell(2, gene_col).value == "BRCA1"

    def test_excluded_column_no_for_retained(self, tmp_path):
        v = _make_variant(acmg_classification="VUS")
        out = str(tmp_path / "out.xlsx")
        write_spreadsheet([v], out, {"sample_id": "S1"})
        wb = openpyxl.load_workbook(out)
        ws = wb["VUS"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        excl_col = headers.index("Excluded") + 1
        assert ws.cell(2, excl_col).value == "No"

    def test_sample_id_in_row(self, tmp_path):
        v = _make_variant(acmg_classification="VUS")
        out = str(tmp_path / "out.xlsx")
        write_spreadsheet([v], out, {"sample_id": "MYSAMPLE"})
        wb = openpyxl.load_workbook(out)
        ws = wb["VUS"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        sid_col = headers.index("Sample ID") + 1
        assert ws.cell(2, sid_col).value == "MYSAMPLE"


# ---------------------------------------------------------------------------
# Disclaimer text integrity
# ---------------------------------------------------------------------------

class TestDisclaimerIntegrity:
    def test_mandatory_disclaimer_exact_text(self):
        assert MANDATORY_DISCLAIMER == (
            "Top genomic candidate(s) requiring expert review. "
            "Variant pathogenicity alone does not prove cause of death."
        )

    def test_acmg_colours_defined_for_all_classes(self):
        for cls in ("Pathogenic", "Likely Pathogenic", "VUS", "Likely Benign", "Benign"):
            assert cls in _ACMG_COLOURS
